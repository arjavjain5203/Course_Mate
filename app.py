import os
import uuid
import shutil
import time
import gc
import base64
import json
import asyncio
import logging
from collections import defaultdict
from datetime import datetime

from fastapi import FastAPI, UploadFile, File, BackgroundTasks, HTTPException, Depends, Header, Request, Form
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse, FileResponse
from dotenv import load_dotenv

from sqlalchemy import create_engine, Column, String, Boolean, DateTime, Integer, Text, ForeignKey
from sqlalchemy.orm import declarative_base, sessionmaker, relationship

from langchain_core.documents import Document
from langchain_mistralai import ChatMistralAI
from langchain_chroma import Chroma
from langchain_core.prompts import ChatPromptTemplate, MessagesPlaceholder
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_huggingface.embeddings import HuggingFaceEndpointEmbeddings
from langchain_google_genai import ChatGoogleGenerativeAI
from langchain_core.messages import HumanMessage, AIMessage

import fitz  # PyMuPDF
import firebase_admin
from firebase_admin import credentials, auth

load_dotenv()

# Setup logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# ==========================================
# 1. Firebase Initialization (Aggressive Diagnostic Mode)
# ==========================================
firebase_cred_path = os.getenv("FIREBASE_CREDENTIALS_PATH")
firebase_cred_json = os.getenv("FIREBASE_CREDENTIALS_JSON")

try:
    firebase_admin.get_app()
except ValueError:
    if firebase_cred_json:
        try:
            cred = credentials.Certificate(json.loads(firebase_cred_json))
            firebase_admin.initialize_app(cred)
            logger.info("✅ [FIREBASE INIT SUCCESS] Admin SDK initialized from FIREBASE_CREDENTIALS_JSON.")
        except Exception as e:
            logger.error(f"❌ [FIREBASE INIT ERROR] FIREBASE_CREDENTIALS_JSON could not be loaded. Details: {e}")
            raise RuntimeError("Failed to initialize Firebase Admin SDK from FIREBASE_CREDENTIALS_JSON") from e
    elif firebase_cred_path:
        abs_path = os.path.abspath(firebase_cred_path)
        if os.path.exists(abs_path):
            try:
                cred = credentials.Certificate(abs_path)
                firebase_admin.initialize_app(cred)
                logger.info(f"✅ [FIREBASE INIT SUCCESS] Admin SDK securely initialized via absolute path: {abs_path}")
            except Exception as e:
                logger.error(f"❌ [FIREBASE INIT ERROR] Credentials file found at {abs_path}, but initialization failed. Details: {e}")
                raise RuntimeError(f"Failed to initialize Firebase Admin SDK using the credentials file at {abs_path}") from e
        else:
            logger.error(f"❌ [FIREBASE CRITICAL WARNING] FIREBASE_CREDENTIALS_PATH was specified as '{firebase_cred_path}'.")
            logger.error(f"❌ Resolved Absolute Search Path: {abs_path}")
            logger.error("❌ NO FILE EXISTS AT THIS EXACT LOCATION. All authenticated endpoints will return 401.")
    else:
        logger.warning("⚠️ [FIREBASE WARNING] FIREBASE_CREDENTIALS_PATH is missing from your .env file.")
        logger.warning("⚠️ Attempting implicit default credentials...")
        try:
            firebase_admin.initialize_app()
            logger.info("✅ [FIREBASE INIT SUCCESS] Initialized successfully via application default credentials.")
        except Exception as e:
            logger.error(f"❌ [FIREBASE INIT ERROR] Default initialization failed: {e}. Auth checks will fail.")


# ==========================================
# 2. Database Schema & PaaS Resiliency
# ==========================================
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_DIR = os.path.join(BASE_DIR, "db")
TEMP_DIR = os.path.join(BASE_DIR, "temp")

os.makedirs(DB_DIR, exist_ok=True)
os.makedirs(TEMP_DIR, exist_ok=True)

# Dialect sanitization & Pooling
DATABASE_URL = os.getenv("DATABASE_URL")
if DATABASE_URL:
    if DATABASE_URL.startswith("postgres://"):
        DATABASE_URL = DATABASE_URL.replace("postgres://", "postgresql://", 1)
    
    engine_args = {}
    if "postgresql" in DATABASE_URL:
        engine_args.update({
            "pool_size": 20,
            "max_overflow": 10,
            "pool_pre_ping": True
        })
    engine = create_engine(DATABASE_URL, **engine_args)
else:
    SQLITE_URL = f"sqlite:///{os.path.join(DB_DIR, 'app.db')}"
    engine = create_engine(SQLITE_URL, connect_args={"check_same_thread": False})

SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
Base = declarative_base()

class DBDocument(Base):
    __tablename__ = "documents"
    id = Column(String, primary_key=True, index=True) # session_id
    user_id = Column(String, index=True, nullable=False)
    filename = Column(String, nullable=False)
    uploaded_at = Column(DateTime, default=datetime.utcnow)
    is_deleted_from_disk = Column(Boolean, default=False)
    
    chats = relationship("DBChat", back_populates="document", cascade="all, delete-orphan")

class DBChat(Base):
    __tablename__ = "chats"
    id = Column(Integer, primary_key=True, autoincrement=True)
    document_id = Column(String, ForeignKey("documents.id"))
    role = Column(String, nullable=False) # 'user' or 'ai'
    message = Column(Text, nullable=False)
    timestamp = Column(DateTime, default=datetime.utcnow)
    
    document = relationship("DBDocument", back_populates="chats")

Base.metadata.create_all(bind=engine)

def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()

# ==========================================
# 3. Authentication Dependency
# ==========================================
def verify_firebase_token(authorization: str = Header(None)):
    logger.debug("\n--- [AUTH DEBUG SEQUENCE START] ---")
    if not authorization:
        logger.error("❌ [AUTH ERROR] The 'Authorization' header is entirely missing from the incoming request.")
        raise HTTPException(status_code=401, detail="Missing Authorization header")
    
    if not authorization.startswith("Bearer "):
        logger.error(f"❌ [AUTH ERROR] Malformed header. Expected 'Bearer <token>'. Received prefix: {authorization[:15]}...")
        raise HTTPException(status_code=401, detail="Invalid Authorization header format. Must start with 'Bearer '")
    
    token = authorization.split("Bearer ")[1]
    logger.debug(f"🔍 [AUTH CHECK] Token extraction successful. Parsing JWT string (Length: {len(token)})...")
        
    try:
        # Granular introspection catch
        # Allows a 10-second grace period for clock differences
        decoded_token = auth.verify_id_token(token, clock_skew_seconds=10)
        logger.debug(f"✅ [AUTH SUCCESS] Token mathematically validated. Authorized UID: {decoded_token.get('uid')}")
        logger.debug("--- [AUTH DEBUG SEQUENCE END] ---\n")
        return decoded_token["uid"]
    except Exception as e:
        logger.error(f"❌ [AUTH CRITICAL EXCEPTION] verify_id_token() rejected the token!")
        logger.error(f"❌ [FIREBASE EXCEPTION DETAILS]: {str(e)}")
        logger.debug("--- [AUTH DEBUG SEQUENCE END] ---\n")
        raise HTTPException(status_code=401, detail=f"Invalid or expired Firebase token. Error: {str(e)}")

# ==========================================
# 4. App Initialization & AI Models
# ==========================================
app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"], 
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

embedding_model = HuggingFaceEndpointEmbeddings(
    model="sentence-transformers/all-MiniLM-L6-v2",
    huggingfacehub_api_token=os.getenv("HUGGINGFACEHUB_API_TOKEN")
)
llm = ChatMistralAI(
    model="mistral-small-latest", 
    temperature=0,
    api_key=os.getenv("MISTRAL_API_KEY")
)

gemini_ocr_llm = ChatGoogleGenerativeAI(
    model="gemini-2.5-flash", 
    temperature=0,
    api_key=os.getenv("GOOGLE_API_KEY")
)
mistral_ocr_llm = ChatMistralAI(
    model="pixtral-12b-2409", 
    temperature=0, 
    timeout=45,
    api_key=os.getenv("MISTRAL_API_KEY")
)

MAX_FILE_SIZE_MB = 10
MAX_PAGE_LIMIT = 10

processing_states = {}
session_locks = defaultdict(asyncio.Lock)
active_vector_stores = {}

# Prompt with Blended Context & Conversational Memory
prompt = ChatPromptTemplate.from_messages([
    ("system", (
        "You are an expert AI study assistant. Follow these strict rules:\n"
        "1. Prioritize answering using the provided 'Document Context'.\n"
        "2. If the user asks a direct question about facts in the document, and the context doesn't contain it, state that the document does not contain this information.\n"
        "3. **Blended Context Rule:** If the user asks a follow-up question, requests an explanation, asks for a comparison, or asks general knowledge questions, you ARE ALLOWED to use your general outside knowledge to supplement your answer.\n"
        "4. If using outside knowledge, clearly distinguish it from document facts.\n"
        "5. Use Markdown for formatting.\n"
        "\nDocument Context:\n{context}"
    )),
    MessagesPlaceholder(variable_name="chat_history"),
    ("human", "{question}")
])

def update_state(session_id: str, status: str, progress: int, message: str, is_error: bool = False):
    processing_states[session_id] = {
        "status": status,
        "progress": progress,
        "message": message,
        "is_error": is_error,
        "timestamp": time.time()
    }

async def scheduled_cleanup(session_id: str):
    """Async cleanup task — runs after a 10-minute grace period."""
    logger.info(f"⏳ Cleanup timer started for session: {session_id} (10 minutes)")
    await asyncio.sleep(600)

    db_path = os.path.join(DB_DIR, session_id)
    temp_file = os.path.join(TEMP_DIR, f"{session_id}.pdf")

    async with session_locks[session_id]:
        gc.collect()

        db = SessionLocal()
        try:
            doc = db.query(DBDocument).filter(DBDocument.id == session_id).first()
            if doc:
                doc.is_deleted_from_disk = True
                db.commit()
        except Exception as e:
            logger.error(f"❌ DB Error in cleanup: {e}")
        finally:
            db.close()

        # Remove from active cache
        if session_id in active_vector_stores:
            del active_vector_stores[session_id]

        if os.path.exists(db_path):
            for _ in range(5):
                try:
                    shutil.rmtree(db_path)
                    logger.info(f"✅ SUCCESS: Deleted vectorstore folder for {session_id}")
                    break
                except PermissionError:
                    await asyncio.sleep(2)
                except Exception as e:
                    logger.error(f"❌ ERROR deleting vectorstore: {e}")
                    break

        if os.path.exists(temp_file):
            try:
                os.remove(temp_file)
                logger.info(f"✅ SUCCESS: Deleted PDF file for {session_id}")
            except Exception as e:
                logger.error(f"❌ ERROR deleting PDF: {e}")

        if session_id in processing_states:
            del processing_states[session_id]
            
    # Cleanup lock after completion
    if session_id in session_locks:
        del session_locks[session_id]

    logger.info(f"🏁 Finished cleanup routine for {session_id}")


async def run_pipeline(session_id: str, upload_path: str, db_path: str, filename: str, filename_lower: str):
    """Heavy async pipeline: OCR extraction → text splitting → Chroma vectorization."""
    try:
        update_state(session_id, "starting", 5, "Starting initialization protocols...")

        cleaned_docs = []

        if filename_lower.endswith(('.doc', '.docx')):
            update_state(session_id, "parsing", 20, "Parsing DOCX structural layout...")
            try:
                import docx as _docx
                doc_struct = _docx.Document(upload_path)
                combined_text = "\n".join([p.text for p in doc_struct.paragraphs if p.text.strip()])
                if not combined_text.strip():
                    update_state(session_id, "failed", 0, "The uploaded document appears to contain no readable text. Please check the file contents.", is_error=True)
                    if os.path.exists(upload_path): os.remove(upload_path)
                    return
                cleaned_docs.append(Document(page_content=combined_text, metadata={"source": filename, "page": 1}))
            except Exception as e:
                update_state(session_id, "failed", 0, f"Failed to parse DOCX structure. Ensure the file is a valid Word document. Detail: {e}", is_error=True)
                if os.path.exists(upload_path): os.remove(upload_path)
                return
        else:
            try:
                doc = fitz.open(upload_path)
            except Exception as e:
                update_state(session_id, "failed", 0, f"Failed to read PDF structure: {e}", is_error=True)
                if os.path.exists(upload_path): os.remove(upload_path)
                return

            num_pages = len(doc)
            requires_ocr = False
            native_texts = []

            for i in range(num_pages):
                text = " ".join(doc[i].get_text().split()).strip()
                native_texts.append(text)
                if len(text) < 40:
                    requires_ocr = True

            active_ocr_engine = "gemini"

            if requires_ocr:
                update_state(session_id, "ocr_preflight", 15, "Scanned layout matrix detected. Activating OCR Preflight Diagnostics Engine...")
                try:
                    await asyncio.get_event_loop().run_in_executor(None, lambda: gemini_ocr_llm.invoke("ping"))
                    update_state(session_id, "ocr_processing", 25, "Primary Vision Core (Gemini) passed status diagnostics. Executing extraction blocks...")
                except Exception as e:
                    if "429" in str(e) or "RESOURCE_EXHAUSTED" in str(e) or "Quota" in str(e):
                        active_ocr_engine = "mistral"
                        update_state(session_id, "fallback_active", 30, "Gemini Exhaustion Rule Triggered (429 Rate Limit). Instantly routing workflows to Mistral-OCR Cluster Core...")
                        try:
                            await asyncio.get_event_loop().run_in_executor(None, lambda: mistral_ocr_llm.invoke("ping"))
                        except Exception as mistral_err:
                            update_state(session_id, "failed", 0, f"OCR Critical Fault: Both extraction engines are unresponsive. Mistral Log: {mistral_err}", is_error=True)
                            doc.close()
                            if os.path.exists(upload_path): os.remove(upload_path)
                            return
                    else:
                        update_state(session_id, "failed", 0, f"Primary Core Allocation Error: {str(e)}", is_error=True)
                        doc.close()
                        if os.path.exists(upload_path): os.remove(upload_path)
                        return
            else:
                update_state(session_id, "parsing", 40, "Standard layout configuration confirmed. Running recursive context parsing matrix layers...")

            OCR_PAGE_TIMEOUT = 90  

            for i in range(num_pages):
                text = native_texts[i]

                if len(text) < 40:
                    pix = doc[i].get_pixmap(dpi=150)
                    img_bytes = pix.tobytes("png")
                    img_base64 = base64.b64encode(img_bytes).decode("utf-8")

                    human_payload = HumanMessage(
                        content=[
                            {"type": "text", "text": "Extract all readable text from this document page. Return only extracted text exactly as written without explanations."},
                            {"type": "image_url", "image_url": {"url": f"data:image/png;base64,{img_base64}"}}
                        ]
                    )

                    pct = int(40 + ((i + 1) / num_pages) * 35)
                    ocr_text = ""
                    page_failed = False

                    await asyncio.sleep(0.5)

                    try:
                        if active_ocr_engine == "gemini":
                            update_state(session_id, "ocr_processing", pct, f"Processing page layer {i+1} of {num_pages} via Gemini Multimodal Vision matrix...")
                            response = await asyncio.wait_for(
                                asyncio.get_event_loop().run_in_executor(None, lambda: gemini_ocr_llm.invoke([human_payload])),
                                timeout=OCR_PAGE_TIMEOUT
                            )
                            ocr_text = response.content.strip()
                        else:
                            update_state(session_id, "fallback_active", pct, f"Processing page layer {i+1} of {num_pages} via Mistral Vision Backup cluster...")
                            response = await asyncio.wait_for(
                                asyncio.get_event_loop().run_in_executor(None, lambda: mistral_ocr_llm.invoke([human_payload])),
                                timeout=OCR_PAGE_TIMEOUT
                            )
                            ocr_text = response.content.strip()

                    except (asyncio.TimeoutError, Exception) as e:
                        error_msg = str(e)
                        is_timeout = isinstance(e, asyncio.TimeoutError)

                        if active_ocr_engine == "gemini" and not is_timeout and \
                                ("429" in error_msg or "RESOURCE_EXHAUSTED" in error_msg or "Quota" in error_msg):
                            active_ocr_engine = "mistral"
                            update_state(session_id, "fallback_active", pct, f"Cascade Error on page {i+1}! Shifting queue to Mistral Core...")
                            try:
                                response = await asyncio.wait_for(
                                    asyncio.get_event_loop().run_in_executor(None, lambda: mistral_ocr_llm.invoke([human_payload])),
                                    timeout=OCR_PAGE_TIMEOUT
                                )
                                ocr_text = response.content.strip()
                            except (asyncio.TimeoutError, Exception) as fallback_err:
                                fb_msg = "Request timed out" if isinstance(fallback_err, asyncio.TimeoutError) else str(fallback_err)
                                page_failed = True
                                update_state(session_id, "failed", 0, f"Network Pipeline Timeout: Both OCR engines failed on page {i+1}. Mistral error — {fb_msg}", is_error=True)
                        else:
                            label = "Request timed out after 90s" if is_timeout else error_msg
                            engine_name = "Gemini" if active_ocr_engine == "gemini" else "Mistral"
                            page_failed = True
                            update_state(session_id, "failed", 0, f"Network Pipeline Timeout: {engine_name} OCR dropped on page {i+1}. Detail — {label}", is_error=True)

                    if page_failed:
                        doc.close()
                        if os.path.exists(upload_path):
                            try: os.remove(upload_path)
                            except: pass
                        return

                    if ocr_text:
                        cleaned_docs.append(Document(page_content=ocr_text, metadata={"source": filename, "page": i+1}))
                else:
                    cleaned_docs.append(Document(page_content=text, metadata={"source": filename, "page": i+1}))

            doc.close()

        if not cleaned_docs or not cleaned_docs[0].page_content.strip():
            update_state(session_id, "failed", 0, "Parsing Pipeline Timeout: No readable character strings or image blocks could be synthesized.", is_error=True)
            if os.path.exists(upload_path): os.remove(upload_path)
            return

        update_state(session_id, "indexing", 85, "Structuring context chunk distributions and vector maps split rules...")
        splitter = RecursiveCharacterTextSplitter(chunk_size=800, chunk_overlap=100)
        chunks = splitter.split_documents(cleaned_docs)

        def _build_vectorstore():
            vs = Chroma.from_documents(
                documents=chunks,
                embedding=embedding_model,
                persist_directory=db_path
            )
            del vs

        await asyncio.get_event_loop().run_in_executor(None, _build_vectorstore)

        update_state(session_id, "completed", 100, "Analysis complete! Content blocks processed into localized vector store indexes.")
        
        asyncio.create_task(scheduled_cleanup(session_id))

    except Exception as e:
        update_state(session_id, "failed", 0, f"Unexpected pipeline fault: {str(e)}", is_error=True)
        if os.path.exists(upload_path): 
            try: os.remove(upload_path)
            except: pass


# ==========================================
# 5. API Endpoints
# ==========================================

@app.get("/")
async def root():
    html_path = os.path.join(BASE_DIR, "app.html")
    if os.path.exists(html_path):
        return FileResponse(html_path)
    return {"status": "Backend is running, but app.html is missing."}


@app.get("/health")
async def health_check():
    return {"status": "ok", "service": "CourseMate AI"}

@app.get("/progress/{session_id}")
async def get_progress_stream(session_id: str):
    async def event_generator():
        while True:
            if session_id in processing_states:
                state = processing_states[session_id]
                payload = {
                    "status": state["status"],
                    "progress": state["progress"],
                    "message": state["message"],
                    "is_error": state["is_error"]
                }
                yield f"data: {json.dumps(payload)}\n\n"

                if state['progress'] == 100 or state['is_error']:
                    break
            else:
                fallback_payload = {
                    "status": "initializing",
                    "progress": 0,
                    "message": "Waiting for stream connection handshake...",
                    "is_error": False
                }
                yield f"data: {json.dumps(fallback_payload)}\n\n"
            await asyncio.sleep(0.5)

    return StreamingResponse(event_generator(), media_type="text/event-stream")


@app.post("/upload")
async def upload_pdf(
    background_tasks: BackgroundTasks, 
    file: UploadFile = File(...),
    session_id: str = Form(None),
    user_id: str = Depends(verify_firebase_token),
    db: SessionLocal = Depends(get_db)
):
    is_reupload = False
    
    if session_id:
        existing_doc = db.query(DBDocument).filter(DBDocument.id == session_id, DBDocument.user_id == user_id).first()
        if existing_doc:
            is_reupload = True
            existing_doc.is_deleted_from_disk = False
            db.commit()
        else:
            session_id = str(uuid.uuid4())
    else:
        session_id = str(uuid.uuid4())

    filename_lower = file.filename.lower()
    allowed_extensions = ('.pdf', '.doc', '.docx')
    if not filename_lower.endswith(allowed_extensions):
        return {"error": "Unsupported File Type! Access allowed to PDF and DOC/DOCX files only.", "session_id": session_id}

    try:
        file_bytes = await file.read()
    except Exception as e:
        return {"error": f"Streaming error reading input parameters: {str(e)}", "session_id": session_id}

    file_size_mb = len(file_bytes) / (1024 * 1024)

    if file_size_mb > MAX_FILE_SIZE_MB:
        return {"error": f"File payload validation failure! Current file size ({file_size_mb:.2f}MB) exceeds system limitations of {MAX_FILE_SIZE_MB}MB.", "session_id": session_id}

    if filename_lower.endswith('.pdf'):
        try:
            import io
            doc_check = fitz.open(stream=io.BytesIO(file_bytes), filetype="pdf")
            num_pages_check = len(doc_check)
            doc_check.close()
            if num_pages_check > MAX_PAGE_LIMIT:
                return {"error": f"Document layout boundary rejection! File length ({num_pages_check} pages) exceeds current system limit rules ({MAX_PAGE_LIMIT} pages max).", "session_id": session_id}
        except Exception as e:
            return {"error": f"Failed to read PDF structure during pre-flight: {e}", "session_id": session_id}

    if not is_reupload:
        new_doc = DBDocument(id=session_id, user_id=user_id, filename=file.filename)
        db.add(new_doc)
        db.commit()

    upload_path = os.path.join(TEMP_DIR, f"{session_id}.pdf")
    vector_db_path = os.path.join(DB_DIR, session_id)

    with open(upload_path, "wb") as f:
        f.write(file_bytes)

    update_state(session_id, "starting", 1, "Starting initialization protocols...")

    background_tasks.add_task(run_pipeline, session_id, upload_path, vector_db_path, file.filename, filename_lower)

    return {"session_id": session_id, "filename": file.filename}


@app.post("/chat")
async def chat(
    data: dict,
    user_id: str = Depends(verify_firebase_token),
    db: SessionLocal = Depends(get_db)
):
    query = data.get("query")
    session_id = data.get("session_id")
    vector_db_path = os.path.join(DB_DIR, session_id)

    if not session_id:
        return {"answer": "Invalid session."}

    doc = db.query(DBDocument).filter(DBDocument.id == session_id, DBDocument.user_id == user_id).first()
    if not doc:
        return {"answer": "Document not found or access denied."}

    async with session_locks[session_id]:
        # Refreshed state inside lock in case it was deleted
        db.refresh(doc)
        if doc.is_deleted_from_disk or not os.path.exists(vector_db_path):
            return {
                "answer": "🔒 For your privacy, the original document was automatically deleted from our servers. I remember our conversation, but to ask new document-specific questions, please re-upload the PDF to this chat.",
                "requires_reupload": True
            }

        past_chats = db.query(DBChat).filter(DBChat.document_id == session_id).order_by(DBChat.timestamp.asc()).all()
        chat_history = []
        
        for msg in past_chats[-6:]:
            if msg.role == 'user':
                chat_history.append(HumanMessage(content=msg.message))
            else:
                chat_history.append(AIMessage(content=msg.message))

        # Vector Store In-Memory Cache Initialization
        if session_id not in active_vector_stores:
            active_vector_stores[session_id] = Chroma(
                persist_directory=vector_db_path,
                embedding_function=embedding_model
            )
        vectorstore = active_vector_stores[session_id]
        
        retriever = vectorstore.as_retriever(
            search_type="mmr", 
            search_kwargs={"k": 3, "fetch_k": 10}
        )
        
        try:
            docs = retriever.invoke(query)
        except Exception as e:
            logger.error(f"❌ [VECTOR STORE RETRIEVAL ERROR] Failed to fetch context for session {session_id}.", exc_info=True)
            raise HTTPException(status_code=500, detail="Vector store I/O operation failed. Please try your request again.")

        context = "No relevant context found in document."
        if docs:
            current_time_str = datetime.now().strftime("%A, %B %d, %Y")
            context_elements = [f"Current Live System Date/Time: {current_time_str}"]
            context_elements.extend(d.page_content for d in docs)
            context = "\n\n".join(context_elements)

        final_prompt = prompt.invoke({
            "context": context, 
            "chat_history": chat_history,
            "question": query
        })
        
        response = llm.invoke(final_prompt)
        answer_text = response.content

        user_msg = DBChat(document_id=session_id, role='user', message=query)
        ai_msg = DBChat(document_id=session_id, role='ai', message=answer_text)
        db.add(user_msg)
        db.add(ai_msg)
        db.commit()

        # Cache remains alive, let python GC handle unused memory. Do not delete vectorstore manually here.

    return {"answer": answer_text}


@app.get("/history")
async def get_history(
    user_id: str = Depends(verify_firebase_token),
    db: SessionLocal = Depends(get_db)
):
    """Fetch user's document history for sidebar."""
    docs = db.query(DBDocument).filter(DBDocument.user_id == user_id).order_by(DBDocument.uploaded_at.desc()).all()
    
    result = []
    for doc in docs:
        # Check physical presence since server might have restarted
        physical_path = os.path.join(DB_DIR, doc.id)
        if not doc.is_deleted_from_disk and not os.path.exists(physical_path):
            doc.is_deleted_from_disk = True
            db.commit()
            
        result.append({
            "session_id": doc.id,
            "filename": doc.filename,
            "uploaded_at": doc.uploaded_at.isoformat(),
            "is_expired": doc.is_deleted_from_disk
        })
        
    return {"history": result}


@app.get("/chat/{session_id}")
async def get_chat_history(
    session_id: str,
    user_id: str = Depends(verify_firebase_token),
    db: SessionLocal = Depends(get_db)
):
    """Fetch previous messages for a specific session."""
    doc = db.query(DBDocument).filter(DBDocument.id == session_id, DBDocument.user_id == user_id).first()
    if not doc:
        raise HTTPException(status_code=404, detail="Session not found")
        
    chats = db.query(DBChat).filter(DBChat.document_id == session_id).order_by(DBChat.timestamp.asc()).all()
    
    messages = []
    for chat in chats:
        messages.append({
            "role": chat.role,
            "message": chat.message,
            "timestamp": chat.timestamp.isoformat()
        })
        
    return {
        "filename": doc.filename,
        "is_expired": doc.is_deleted_from_disk,
        "messages": messages
    }


@app.get("/storage-status")
async def get_storage_status():
    temp_files = os.listdir(TEMP_DIR) if os.path.exists(TEMP_DIR) else []
    db_folders = os.listdir(DB_DIR) if os.path.exists(DB_DIR) else []

    return {
        "active_pdfs_count": len(temp_files),
        "active_pdfs": temp_files,
        "active_db_sessions_count": len(db_folders),
        "active_db_sessions": db_folders,
        "location_info": {
            "pdf_storage": os.path.abspath(TEMP_DIR),
            "vector_storage": os.path.abspath(DB_DIR)
        }
    }
